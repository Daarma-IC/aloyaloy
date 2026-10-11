"""
Template Excel pengukuran QoS LoRa Meshta (RSSI/SNR/PSP/latensi vs jarak per SF).

Alur: uji di lapangan pakai "Uji QoS" di aplikasi → ekspor CSV ringkasan → tempel ke sheet CSV_Ringkasan →
ketik run_id di sheet Pengukuran (baris SF & jarak yang sesuai) → Rekap & grafik terisi otomatis.
Tanpa aplikasi: ketik angka langsung di kolom otomatis (menimpa rumusnya).

  pip install openpyxl && python make_sheet.py ~/Downloads/Meshta-Pengukuran-QoS.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

SFS = [7, 8, 9, 10, 11, 12]
DISTANCES = [100, 250, 500, 750, 1000, 1500, 2000, 3000, 4000, 5000]
# Sensitivitas & batas SNR demodulasi SX1276 @125 kHz (datasheet Semtech, tabel 13 & 14).
SENSITIVITY = {7: -123, 8: -126, 9: -129, 10: -132, 11: -134.5, 12: -137}
SNR_LIMIT = {7: -7.5, 8: -10, 9: -12.5, 10: -15, 11: -17.5, 12: -20}
SUMMARY_HEADER = ["run_id", "label", "sender", "sf", "total_sent", "received", "psp_percent", "mean_rssi_dbm",
                  "mean_snr_db", "mean_distance_m", "mean_latency_ms", "started_at_utc"]
PACKET_HEADER = ["run_id", "label", "sender", "seq", "total", "received_at_utc", "rssi_dbm", "snr_db", "sf", "bw_khz",
                 "hop_left", "meta_matched", "ble_rssi_dbm", "distance_m", "sender_position_age_s", "latency_ms", "path"]

BLUE = "1F5FAF"
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BLUE)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")   # kuning muda = diisi manusia
AUTO_FILL = PatternFill("solid", fgColor="EAF2FB")    # biru muda = rumus otomatis
THIN = Side(style="thin", color="C9D3DF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row, i, t)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, CENTER, BOX
        if widths:
            ws.column_dimensions[col(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 32


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="5B6B7F")


wb = Workbook()

# ---------------------------------------------------------------- Petunjuk
ws = wb.active
ws.title = "Petunjuk"
ws.column_dimensions["A"].width = 110
lines = [
    ("Template Pengukuran QoS LoRa — Meshta / Nusa Node", Font(bold=True, size=14, color=BLUE)),
    ("Parameter: RSSI, SNR, PSP (Packet Success Percentage), PLR, latensi terhadap jarak, untuk SF7–SF12.", None),
    ("", None),
    ("ALUR DENGAN APLIKASI", Font(bold=True)),
    ("1. Flash semua node dengan firmware SF yang sama (firmware/sf-variants/NusaNode_SFx).", None),
    ("2. Di titik uji: buka chat → Uji QoS, beri label (mis. 'SF7 500m LOS'), kirim N paket (disarankan 50–100).", None),
    ("3. Penerima: menu QoS → Ekspor CSV ringkasan. Buka CSV, salin semua baris, tempel ke sheet CSV_Ringkasan sel A1.", None),
    ("   (Opsional: CSV per paket ditempel ke CSV_Paket untuk lampiran/sebaran data.)", None),
    ("4. Di sheet Pengukuran: cari baris SF & jarak target yang sesuai, ketik run_id di kolom kuning 'run_id'.", None),
    ("   Kolom biru (jarak GPS, terkirim, diterima, RSSI, SNR, latensi) terisi otomatis; PSP & PLR dihitung.", None),
    ("5. Sheet Rekap: tabel rata-rata per SF × jarak + grafik RSSI, SNR, PSP vs jarak. Siap dipakai di bab hasil.", None),
    ("", None),
    ("TANPA APLIKASI / DATA MANUAL", Font(bold=True)),
    ("Ketik angka langsung di kolom biru (menimpa rumus). Rekap & grafik tetap jalan.", None),
    ("", None),
    ("CATATAN", Font(bold=True)),
    ("• Kuning = diisi saat pengukuran; biru = otomatis. Baris SF/jarak boleh ditambah di bawah (salin baris terakhir).", None),
    ("• Jarak di Rekap dikelompokkan menurut 'Jarak target'; jarak GPS sebenarnya tetap tercatat untuk laporan.", None),
    ("• Ulangi pengukuran yang sama ≥3 kali (Ulangan 1,2,3) bila sempat; Rekap otomatis merata-rata semua ulangan.", None),
    ("• Catat kondisi (LOS/NLOS), tinggi antena, cuaca — penguji biasanya menanyakan ini.", None),
    ("• Sheet Teori: Time-on-Air & sensitivitas per SF (pembanding hasil ukur). Ubah ukuran payload di sel kuning.", None),
    ("• PSP = diterima ÷ terkirim × 100%. PLR = 100% − PSP. RSSI di bawah garis sensitivitas → paket mulai hilang.", None),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(i, 1, t).font = f or Font()

# ---------------------------------------------------------------- CSV sheets (tempel dari aplikasi)
csv1 = wb.create_sheet("CSV_Ringkasan")
header(csv1, 1, SUMMARY_HEADER, [14, 18, 14, 6, 10, 10, 11, 13, 12, 14, 14, 22])
csv1.freeze_panes = "A2"
csv2 = wb.create_sheet("CSV_Paket")
header(csv2, 1, PACKET_HEADER, [14, 18, 14, 6, 7, 22, 9, 8, 5, 7, 8, 11, 12, 11, 12, 11, 30])
csv2.freeze_panes = "A2"

# ---------------------------------------------------------------- Pengukuran
pm = wb.create_sheet("Pengukuran", 1)
title(pm, "Lembar Pengukuran", "Kuning = isi di lapangan · Biru = otomatis dari CSV_Ringkasan (boleh ditimpa manual)")
cols = ["No", "SF", "Jarak target (m)", "Ulangan", "Lokasi / titik", "Kondisi", "Tinggi antena TX (m)",
        "Tinggi antena RX (m)", "Tanggal", "Cuaca", "run_id", "Jarak GPS (m)", "Paket terkirim", "Paket diterima",
        "PSP (%)", "PLR (%)", "RSSI rata-rata (dBm)", "SNR rata-rata (dB)", "Latensi rata-rata (ms)", "Catatan"]
widths = [5, 5, 10, 8, 20, 9, 9, 9, 11, 10, 14, 10, 9, 9, 8, 8, 11, 10, 11, 28]
HR = 4
header(pm, HR, cols, widths)
pm.freeze_panes = pm.cell(HR + 1, 3)
kondisi = DataValidation(type="list", formula1='"LOS,NLOS,Semi-LOS"', allow_blank=True)
sf_dv = DataValidation(type="list", formula1='"7,8,9,10,11,12"', allow_blank=True)
pm.add_data_validation(kondisi)
pm.add_data_validation(sf_dv)

# kolom otomatis → kolom di CSV_Ringkasan (rentang terbatas, bukan A:A, supaya hitung ulang ringan)
CSV_ROWS = 2000
RUN_IDS = f"CSV_Ringkasan!$A$2:$A${CSV_ROWS}"
LOOKUP = {12: "J", 13: "E", 14: "F", 17: "H", 18: "I", 19: "K"}
ROWS = 300
for i in range(ROWS):
    r = HR + 1 + i
    planned = i < len(SFS) * len(DISTANCES)
    pm.cell(r, 1, i + 1)
    if planned:
        pm.cell(r, 2, SFS[i // len(DISTANCES)])
        pm.cell(r, 3, DISTANCES[i % len(DISTANCES)])
        pm.cell(r, 4, 1)
    for c in range(2, 12):
        pm.cell(r, c).fill = INPUT_FILL
    pm.cell(r, 20).fill = INPUT_FILL
    for c, src in LOOKUP.items():
        pm.cell(r, c, f'=IF($K{r}="","",IFERROR(INDEX(CSV_Ringkasan!${src}$2:${src}${CSV_ROWS},MATCH($K{r},{RUN_IDS},0)),""))')
        pm.cell(r, c).fill = AUTO_FILL
    pm.cell(r, 15, f'=IF(AND(ISNUMBER(M{r}),ISNUMBER(N{r}),N(M{r})>0),N{r}/M{r}*100,"")').fill = AUTO_FILL
    pm.cell(r, 16, f'=IF(ISNUMBER(O{r}),100-O{r},"")').fill = AUTO_FILL
    pm.cell(r, 9).number_format = "dd/mm/yyyy"
    for c in (12, 15, 16, 17, 18, 19):
        pm.cell(r, c).number_format = "0.0"
    for c in range(1, 21):
        pm.cell(r, c).border = BOX
kondisi.add(f"F{HR + 1}:F{HR + ROWS}")
sf_dv.add(f"B{HR + 1}:B{HR + ROWS}")
# Peringatan: run_id diisi tapi tidak ditemukan di CSV_Ringkasan.
pm.conditional_formatting.add(
    f"K{HR + 1}:K{HR + ROWS}",
    FormulaRule(formula=[f'AND(K{HR + 1}<>"",ISNA(MATCH(K{HR + 1},{RUN_IDS},0)))'],
                fill=PatternFill("solid", fgColor="F8C9C4")),
)
LAST = HR + ROWS

# ---------------------------------------------------------------- Rekap
rk = wb.create_sheet("Rekap", 2)
title(rk, "Rekap per SF × Jarak", "Rata-rata semua ulangan di sheet Pengukuran (dikelompokkan menurut jarak target). Jarak di kolom A boleh diubah.")
rk.column_dimensions["A"].width = 14
METRICS = [("PSP (%)", "O"), ("RSSI rata-rata (dBm)", "Q"), ("SNR rata-rata (dB)", "R"), ("Latensi rata-rata (ms)", "S"),
           ("Jumlah ulangan terisi", None)]
blocks = {}
row = 4
for name, src in METRICS:
    rk.cell(row, 1, name).font = Font(bold=True, color=BLUE, size=12)
    header(rk, row + 1, ["Jarak (m)"] + [f"SF{sf}" for sf in SFS])
    for j, d in enumerate(DISTANCES):
        r = row + 2 + j
        rk.cell(r, 1, d if not blocks else f"=$A${blocks['PSP (%)'] + j}").border = BOX
        rk.cell(r, 1).fill = INPUT_FILL if not blocks else AUTO_FILL
        for k, sf in enumerate(SFS):
            c = rk.cell(r, 2 + k)
            crit = f"Pengukuran!$B${HR + 1}:$B${LAST},{sf},Pengukuran!$C${HR + 1}:$C${LAST},$A{r}"
            if src:
                c.value = f'=IFERROR(AVERAGEIFS(Pengukuran!${src}${HR + 1}:${src}${LAST},{crit}),"")'
                c.number_format = "0.0"
            else:
                c.value = f'=COUNTIFS({crit},Pengukuran!$O${HR + 1}:$O${LAST},">=0")'
            c.border, c.alignment = BOX, CENTER
            c.fill = AUTO_FILL
    blocks[name] = row + 2
    row += len(DISTANCES) + 4
for k in range(len(SFS)):
    rk.column_dimensions[col(2 + k)].width = 10

# Referensi sensitivitas di samping tabel RSSI
ref_col = 2 + len(SFS) + 1
rk.cell(blocks["RSSI rata-rata (dBm)"] - 2, ref_col, "Sensitivitas SX1276 @125 kHz (dBm)").font = Font(bold=True, color=BLUE)
for k, sf in enumerate(SFS):
    rk.cell(blocks["RSSI rata-rata (dBm)"] - 1 + k, ref_col, f"SF{sf}")
    rk.cell(blocks["RSSI rata-rata (dBm)"] - 1 + k, ref_col + 1, SENSITIVITY[sf])
rk.column_dimensions[col(ref_col)].width = 8

# Data grafik tersembunyi: sel kosong → NA() supaya garis tidak jatuh ke 0.
CH = row + 2
rk.cell(CH - 1, 1, "Data grafik (otomatis, jangan diubah)").font = Font(italic=True, color="8896A8")
chart_blocks = {}
crow = CH
for name, _ in METRICS[:4]:
    src_start = blocks[name]
    for j in range(len(DISTANCES)):
        r = crow + j
        rk.cell(r, 1, f"=$A${src_start + j}")
        for k in range(len(SFS)):
            ref = f"{col(2 + k)}{src_start + j}"
            rk.cell(r, 2 + k, f"=IF(ISNUMBER({ref}),{ref},NA())")
    chart_blocks[name] = crow
    crow += len(DISTANCES) + 1
for r in range(CH - 1, crow):
    rk.row_dimensions[r].hidden = True


def chart(name, y_title, anchor, y_min=None, y_max=None):
    ch = ScatterChart()
    ch.title = f"{name.split(' (')[0]} terhadap jarak"
    ch.style = 13
    ch.x_axis.title = "Jarak (m)"
    ch.y_axis.title = y_title
    ch.height, ch.width = 8.5, 16
    ch.visible_cells_only = False   # data grafik ada di baris tersembunyi
    ch.display_blanks = "gap"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if y_min is not None:
        ch.y_axis.scaling.min = y_min
    if y_max is not None:
        ch.y_axis.scaling.max = y_max
    start = chart_blocks[name]
    xs = Reference(rk, min_col=1, min_row=start, max_row=start + len(DISTANCES) - 1)
    for k, sf in enumerate(SFS):
        ys = Reference(rk, min_col=2 + k, min_row=start, max_row=start + len(DISTANCES) - 1)
        s = Series(ys, xs, title=f"SF{sf}")
        s.marker.symbol = "circle"
        s.marker.size = 6
        s.smooth = False
        ch.series.append(s)
    rk.add_chart(ch, anchor)


gx = col(ref_col + 3)
chart("RSSI rata-rata (dBm)", "RSSI (dBm)", f"{gx}4")
chart("PSP (%)", "PSP (%)", f"{gx}22", 0, 100)
chart("SNR rata-rata (dB)", "SNR (dB)", f"{gx}40")
chart("Latensi rata-rata (ms)", "Latensi (ms)", f"{gx}58")

# ---------------------------------------------------------------- Teori
th = wb.create_sheet("Teori")
title(th, "Teori LoRa per SF", "Time-on-Air (rumus Semtech AN1200.13) & sensitivitas SX1276 — pembanding hasil ukur. Ubah sel kuning.")
params = [("Bandwidth (kHz)", 125), ("Coding rate (4/x, isi 5..8)", 5), ("Preamble (simbol)", 8),
          ("Payload (byte)", 60), ("Header eksplisit (1=ya)", 1), ("CRC (1=ya)", 1), ("Daya pancar (dBm)", 14),
          ("Frekuensi (MHz)", 921)]
for i, (k, v) in enumerate(params):
    th.cell(4 + i, 1, k)
    c = th.cell(4 + i, 2, v)
    c.fill, c.border = INPUT_FILL, BOX
th.column_dimensions["A"].width = 28
BW, CR, PRE, PL, EXPL, CRC, PTX = "$B$4", "$B$5", "$B$6", "$B$7", "$B$8", "$B$9", "$B$10"
HT = 14
header(th, HT, ["SF", "T simbol (ms)", "Low DR optimize", "Simbol payload", "Time-on-Air (ms)", "Bitrate (bps)",
                "Sensitivitas (dBm)", "SNR min (dB)", "Link budget (dB)"], [28, 12, 12, 12, 14, 12, 14, 12, 14])
for i, sf in enumerate(SFS):
    r = HT + 1 + i
    th.cell(r, 1, sf)
    th.cell(r, 2, f"=2^A{r}/{BW}")
    th.cell(r, 3, f"=IF(AND(A{r}>=11,{BW}<=125),1,0)")
    th.cell(r, 4, f"=8+MAX(CEILING((8*{PL}-4*A{r}+28+16*{CRC}-20*(1-{EXPL}))/(4*(A{r}-2*C{r})),1)*{CR},0)")
    th.cell(r, 5, f"=({PRE}+4.25)*B{r}+D{r}*B{r}")
    th.cell(r, 6, f"=A{r}*({BW}*1000)/2^A{r}*4/{CR}")
    th.cell(r, 7, SENSITIVITY[sf])
    th.cell(r, 8, SNR_LIMIT[sf])
    th.cell(r, 9, f"={PTX}-G{r}")
    for c in range(1, 10):
        th.cell(r, c).border, th.cell(r, c).alignment = BOX, CENTER
        th.cell(r, c).fill = AUTO_FILL
    for c in (2, 5, 6, 9):
        th.cell(r, c).number_format = "0.0"
th.cell(HT + 8, 1, "Low DR optimize wajib untuk SF11/12 @125 kHz (T simbol > 16 ms). Link budget = daya pancar − sensitivitas (tanpa gain antena).").font = Font(italic=True, color="5B6B7F")

wb.save(sys.argv[1])
print("ok", sys.argv[1])
